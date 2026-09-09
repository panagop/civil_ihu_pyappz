"""Tests for the parts of the περιγράμματα pipeline that need no database.

The Postgres instance is reachable only from inside Railway, so everything that
can be checked without it is checked here: the archive parses, the column spec
and the form agree, the diff does not report phantom changes, and a course
renders into a Word document with every template variable satisfied.
"""

from __future__ import annotations

import io
import sys
from decimal import Decimal
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "streamlit"))

import perigrammata_db as pdb  # noqa: E402
import perigrammata_report as report  # noqa: E402
import seed_perigrammata as seed  # noqa: E402


@pytest.fixture(scope="module")
def courses_2025() -> list[dict]:
    rows, _ = seed.read_archive(
        seed.ARCHIVE_DIR / "perigrammata_gr_2025.csv", 2025, "gr"
    )
    return rows


def test_every_content_column_has_a_form_field():
    """The DDL, the form and the Word context all come from CONTENT_COLUMNS."""
    assert set(pdb.CONTENT_COLUMNS) == set(pdb.FIELD_LABELS)


def test_schema_covers_every_content_column():
    for column in pdb.CONTENT_COLUMNS:
        assert f"{column} " in pdb.SCHEMA_SQL


def test_archive_drops_the_debris_row():
    """The 2018 worksheet has one row with no code and a stray sentence."""
    rows, dropped = seed.read_archive(
        seed.ARCHIVE_DIR / "perigrammata_gr_2018.csv", 2018, "gr"
    )
    assert len(rows) == 102
    assert dropped == 1
    assert all(row["code"] for row in rows)


def test_archive_types(courses_2025):
    assert len(courses_2025) == 96
    for row in courses_2025:
        assert isinstance(row["examino"], int)
        for field in pdb.NUMERIC_COLUMNS:
            value = row[field]
            assert value is None or isinstance(value, float)


def test_numbers_print_without_a_decimal_tail():
    """Page 1 rendered `4.0` into Word where the sheet said `4`."""
    assert report.format_value("hours1", 4.0) == "4"
    assert report.format_value("examino", 1) == "1"
    assert report.format_value("hours1", Decimal("52")) == "52"
    assert report.format_value("hours1", 4.5) == "4,5"
    assert report.format_value("hours1", None) == ""
    assert report.format_value("name", "Στατική") == "Στατική"


def test_docx_context_is_complete(courses_2025):
    from docxtpl import DocxTemplate

    context = report.docx_context(courses_2025[0])
    template = DocxTemplate(io.BytesIO(report._template_bytes("gr")))
    missing = template.get_undeclared_template_variables() - set(context)
    assert missing == set()
    assert all(isinstance(value, str) for value in context.values())


def test_diff_ignores_equivalent_values(courses_2025):
    """Decimal from Postgres and float from the form are not a change."""
    before = dict(courses_2025[0])
    after = dict(before)
    after["hours1"] = Decimal(str(before["hours1"]))
    after["name"] = f"  {before['name']}  "  # the form pads, _normalize strips
    assert pdb.diff_course(before, after) == {}


def test_diff_reports_real_changes(courses_2025):
    before = dict(courses_2025[0])
    after = dict(before, name="Νέος τίτλος", hours1=99.0)
    changes = pdb.diff_course(before, after)
    assert set(changes) == {"name", "hours1"}
    assert changes["name"] == [before["name"], "Νέος τίτλος"]
    assert changes["hours1"][1] == 99.0


def test_blank_text_normalizes_to_none():
    assert pdb._normalize("name", "   ") is None
    assert pdb._normalize("hours1", "") is None
    assert pdb._normalize("hours1", "4,5") == 4.5
    assert pdb._normalize("examino", "3") == 3


def test_render_course_produces_a_document(courses_2025):
    from docx import Document

    data = report.render_course(courses_2025[0], "gr")
    document = Document(io.BytesIO(data))
    text = "\n".join(p.text for cell in document.tables[0]._cells for p in cell.paragraphs)
    assert courses_2025[0]["name"].strip() in text


def test_full_report_merges_every_course(courses_2025):
    from docx import Document

    subset = courses_2025[:3]
    document = Document(io.BytesIO(report.build_full_report(subset, "gr")))
    single = Document(io.BytesIO(report.render_course(subset[0], "gr")))
    assert len(document.tables) >= 3 * len(single.tables)


def test_changes_report_survives_an_empty_period():
    import pandas as pd
    from datetime import datetime

    data = report.build_changes_report(
        pd.DataFrame(), datetime(2026, 1, 1), datetime(2026, 12, 31), curriculum=2025
    )
    assert data[:2] == b"PK"  # a .docx is a zip
